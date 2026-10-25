#!/usr/bin/env python3
"""Convierte un Markdown sencillo en un .docx con formato limpio.

Uso:  python a_docx.py entrada.md salida.docx
Soporta: # títulos (1-3), párrafos, listas con - o 1., **negrita**, tablas | a | b |.
Requiere: pip install python-docx
"""

import re
import sys

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError:
    sys.exit("Falta python-docx. Instálalo con:  pip install python-docx")


def _runs(par, texto):
    for i, trozo in enumerate(re.split(r"\*\*", texto)):
        if trozo:
            par.add_run(trozo).bold = i % 2 == 1


def convertir(md: str, salida: str) -> None:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        l = lineas[i].rstrip()
        if not l.strip():
            i += 1
            continue
        if l.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in celdas):
                    filas.append(celdas)
                i += 1
            ncol = max(len(f) for f in filas)
            t = doc.add_table(rows=len(filas), cols=ncol)
            t.style = "Table Grid"
            for r, fila in enumerate(filas):
                for c, txt in enumerate(fila):
                    p = t.cell(r, c).paragraphs[0]
                    _runs(p, txt)
                    if r == 0:
                        for run in p.runs:
                            run.bold = True
                    if re.fullmatch(r"[\d.,\s€%-]+", txt):
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            continue
        m = re.match(r"(#{1,3})\s+(.*)", l)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
        elif re.match(r"\s*[-*]\s+", l):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", l))
        elif re.match(r"\s*\d+[.)]\s+", l):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", l))
        else:
            _runs(doc.add_paragraph(), l)
        i += 1
    doc.save(salida)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as f:
        convertir(f.read(), sys.argv[2])
    print(f"Creado {sys.argv[2]}")
